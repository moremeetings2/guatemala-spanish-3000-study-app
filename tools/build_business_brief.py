"""Build the short family-facing Word brief from its checked-in source."""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/hablavos-product-and-business-brief.md"
OUTPUT = SOURCE.with_suffix(".docx")

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.85)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.color.rgb = RGBColor(0, 0, 0)

normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
doc.styles["Title"].font.size = Pt(24)
doc.styles["Heading 1"].font.size = Pt(14)
doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
doc.styles["Heading 2"].font.size = Pt(11)

for line in SOURCE.read_text().splitlines():
    if not line:
        continue
    if line.startswith("# "):
        doc.add_paragraph(line[2:], "Title")
    elif line.startswith("## "):
        # Two short pages: learning logic first, business decisions second.
        if line == "## How the free app could earn money":
            doc.add_page_break()
        doc.add_paragraph(line[3:], "Heading 1")
    elif line.startswith("### "):
        doc.add_paragraph(line[4:], "Heading 2")
    elif line.startswith("- "):
        p = doc.add_paragraph(line[2:], "List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.font.size = Pt(9)
    elif line.startswith("For John"):
        doc.add_paragraph(line, "Subtitle")
    else:
        doc.add_paragraph(line)

doc.core_properties.title = "Hablavos product and business brief"
doc.core_properties.subject = "Learning design and proposed free-app revenue strategy"
doc.core_properties.author = "Hablavos"
# The bundled Word template carries a decorative title rule. Keep this brief plain.
for root in (doc.styles.element, doc.element):
    for borders in list(root.iter(qn("w:pBdr"))):
        borders.getparent().remove(borders)
doc.save(OUTPUT)
print(OUTPUT)
