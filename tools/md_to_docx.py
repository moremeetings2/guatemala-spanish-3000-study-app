# -*- coding: utf-8 -*-
"""Markdown -> .docx for the Spanish plan documents.

Targets Google Docs import: real Heading styles, real tables, bullet lists, and
literal-numbered prompt lists. Numbered items are written with their number as
text rather than as Word auto-numbering, because the prompts are referred to by
number ("prompt 7") and must never renumber themselves on import.
"""
import re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE = re.compile(r'(\*\*[^*]+\*\*|(?<!\*)\*[^*\n]+\*(?!\*)|`[^`]+`)')

def add_runs(par, text):
    """Split text on **bold**, *italic* and `code` and add styled runs."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*"):
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)

def hrule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), 'BFBFBF')
    borders.append(bottom); pPr.append(borders)
    p.paragraph_format.space_after = Pt(6)

def flush_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(cols):
            cell = cells[ci]
            cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], row[ci] if ci < len(row) else "")
            if ri == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def convert(md_path, out_path):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    lines = open(md_path, encoding="utf-8").read().split("\n")
    i, table_rows, in_code = 0, [], False

    while i < len(lines):
        raw = lines[i]
        s = raw.strip()

        if s.startswith("```"):
            if in_code:
                in_code = False
            else:
                flush_table(doc, table_rows); table_rows = []
                in_code = True
            i += 1; continue

        if in_code:
            p = doc.add_paragraph()
            r = p.add_run(raw)
            r.font.name = "Consolas"; r.font.size = Pt(9)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Inches(0.2)
            i += 1; continue

        # table rows accumulate until the block ends
        if s.startswith("|") and s.endswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            # a separator row must actually contain dashes; an all-blank row is a
            # fill-in row in the printable templates and must be preserved
            is_sep = any(cells) and all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c)
            if not is_sep:
                table_rows.append(cells)
            i += 1; continue
        if table_rows:
            flush_table(doc, table_rows); table_rows = []

        if not s:
            i += 1; continue

        if re.fullmatch(r'-{3,}', s):
            hrule(doc); i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl = min(len(m.group(1)), 4)
            p = doc.add_paragraph(style="Heading %d" % lvl)
            add_runs(p, m.group(2))
            i += 1; continue

        if s.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            add_runs(p, s[2:])
            for r in p.runs:
                r.italic = True
            i += 1; continue

        # indented numbered prompt (nested under a bullet)
        m = re.match(r'^\s{2,}(\d+)\.\s+(.*)$', raw)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.75)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
            p.add_run("%s. " % m.group(1))
            add_runs(p, m.group(2))
            i += 1; continue

        m = re.match(r'^-\s+(.*)$', s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, m.group(1))
            i += 1; continue

        m = re.match(r'^(\d+)\.\s+(.*)$', s)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(3)
            p.add_run("%s. " % m.group(1))
            add_runs(p, m.group(2))
            i += 1; continue

        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1

    flush_table(doc, table_rows)
    doc.save(out_path)
    return doc

if __name__ == "__main__":
    d = convert(sys.argv[1], sys.argv[2])
    print(sys.argv[2], "paragraphs=%d tables=%d" % (len(d.paragraphs), len(d.tables)))
